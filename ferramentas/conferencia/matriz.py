"""Preenche a Matriz de Atendimento ao Termo de Referência (IN 37/2025) a partir dos textos do projeto.

Uso:
    python -m ferramentas.conferencia.matriz candoi

Para cada item do TR (modelos/orgaos/<orgao>/matrizes/<documento>_<sistema>.yaml), localiza a seção correspondente em
texto/*.md (com a mesma numeração do .docx), verifica pendências [A VERIFICAR], imagens e tabelas
ausentes e classifica: ATENDIDO / PARCIALMENTE ATENDIDO / NÃO ATENDIDO.
Gera saida/Matriz_<Anexo>_<municipio>.xlsx (abas Matriz e Resumo) e imprime o resumo.
"""
import argparse
import re
from collections import Counter
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ferramentas.util import console_utf8, projeto, slug, RAIZ
from ferramentas.docx.montar import carregar_ficha, substituir_variaveis

CORES = {"ATENDIDO": "C6EFCE", "PARCIALMENTE ATENDIDO": "FFEB9C", "NÃO ATENDIDO": "FFC7CE"}
RE_TIT = re.compile(r"^(#{1,4})\s+(.*?)\s*(\{-\})?$")
RE_PEND = re.compile(r"\[A VERIFICAR[^\]]*\]")
RE_ARQ = re.compile(r"^!(?:\[|quadro\[|pdf\[).*?\]\((.+?)\)")


def secoes(texto: str) -> list[dict]:
    """Seções com numeração igual à do montador; 'fim' inclui as subseções."""
    texto = re.sub(r"<!--.*?-->", "", texto, flags=re.S)
    linhas = texto.splitlines()
    cont = [0, 0, 0, 0]
    out = []
    for i, ln in enumerate(linhas):
        m = RE_TIT.match(ln.strip())
        if not m:
            continue
        nivel, titulo = len(m.group(1)), m.group(2)
        numero = ""
        if m.group(3) is None:
            cont[nivel - 1] += 1
            for k in range(nivel, 4):
                cont[k] = 0
            numero = ".".join(str(c) for c in cont[:nivel])
        out.append({"nivel": nivel, "titulo": titulo, "numero": numero, "ini": i})
    for j, s in enumerate(out):
        s["fim"] = next((t["ini"] for t in out[j + 1:] if t["nivel"] <= s["nivel"]), len(linhas))
        s["linhas"] = linhas[s["ini"] + 1:s["fim"]]
    return out


def _casa(titulo: str, padrao: str) -> bool:
    t, p = slug(titulo).split("-"), slug(padrao).split("-")
    return any(t[i:i + len(p)] == p for i in range(len(t) - len(p) + 1))


def avaliar(item: dict, secs: list[dict], base: Path) -> dict:
    achadas = [s for s in secs if any(_casa(s["titulo"], p) for p in item["secao"].split("|"))]
    if not achadas:
        return {"local": "—", "situacao": "NÃO ATENDIDO", "obs": "Seção não encontrada no documento."}
    local = "; ".join(f"{s['numero']} {s['titulo']}".strip() for s in achadas)
    linhas = [ln for s in achadas for ln in s["linhas"] if ln.strip()]
    if item.get("rotulo"):
        rx = re.compile(item["rotulo"], re.I)
        linhas = [ln for ln in linhas if rx.search(ln)]
        if not linhas:
            return {"local": local, "situacao": "NÃO ATENDIDO", "obs": "Informação não localizada na seção."}
    if not linhas:
        return {"local": local, "situacao": "NÃO ATENDIDO", "obs": "Seção sem conteúdo."}
    problemas = []
    for ln in linhas:
        problemas += [p.strip("[]") for p in RE_PEND.findall(ln)]
        m = RE_ARQ.match(ln.strip())
        if m and not (base / m.group(1)).exists():
            problemas.append(f"arquivo ausente: {m.group(1)}")
    problemas = list(dict.fromkeys(problemas))
    if problemas:
        return {"local": local, "situacao": "PARCIALMENTE ATENDIDO", "obs": "; ".join(problemas)[:1500]}
    return {"local": local, "situacao": "ATENDIDO", "obs": ""}


def gerar(cfg) -> Path:
    d = cfg["dir"]
    arq = cfg["orgao_info"]["pasta"] / "matrizes" / f"{cfg['documento']}_{cfg['sistema']}.yaml".lower()
    if not arq.exists():
        raise SystemExit(f"Sem matriz do termo de referência para {arq.stem} no órgão {cfg['orgao']} ({arq}).")
    defin = yaml.safe_load(arq.read_text(encoding="utf-8"))
    ficha = carregar_ficha(d)
    texto = "\n\n".join(f.read_text(encoding="utf-8") for f in sorted((d / "texto").glob("*.md")))
    secs = secoes(substituir_variaveis(texto, ficha))
    resultados = [{**it, **avaliar(it, secs, d)} for it in defin["itens"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz"
    fino = Side(style="thin", color="999999")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    ws.append([f"MATRIZ DE ATENDIMENTO AO {defin['anexo'].upper()} DA {cfg['orgao_info'].get('norma', '').upper()}"])
    ws.append([f"Documento: {cfg['documento']} – {cfg['sistema']} {cfg['municipio']}/{cfg['uf']} – Projeto "
               f"{cfg['empresa'].get('sigla', '')} {cfg.get('numero_projeto') or ''}"])
    ws.append([f"Referência: {defin['titulo']}"])
    ws.append([])
    cab = ["Item", "Informação exigida", "Item correspondente no documento", "Situação", "Pendências / Observação"]
    ws.append(cab)
    for c in ws[5]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="C5E0B3")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda
    ws["A1"].font = Font(bold=True, size=12)
    grupo = None
    for r in resultados:
        if r["grupo"] != grupo:
            grupo = r["grupo"]
            ws.append([grupo])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=5)
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="E2EFDA")
        ws.append([r["item"], r["exigencia"], r["local"], r["situacao"], r["obs"]])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = borda
        ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=CORES[r["situacao"]])
    for col, w in zip("ABCDE", (8, 45, 38, 22, 70)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"

    rs = wb.create_sheet("Resumo")
    rs.append(["Grupo", "Total", "ATENDIDO", "PARCIALMENTE ATENDIDO", "NÃO ATENDIDO"])
    por_grupo = {}
    for r in resultados:
        por_grupo.setdefault(r["grupo"], Counter())[r["situacao"]] += 1
    for g, c in por_grupo.items():
        rs.append([g, sum(c.values()), c["ATENDIDO"], c["PARCIALMENTE ATENDIDO"], c["NÃO ATENDIDO"]])
    tot = Counter(r["situacao"] for r in resultados)
    rs.append(["TOTAL", len(resultados), tot["ATENDIDO"], tot["PARCIALMENTE ATENDIDO"], tot["NÃO ATENDIDO"]])
    for c in rs[1] + rs[rs.max_row]:
        c.font = Font(bold=True)
    rs.column_dimensions["A"].width = 40

    saida = d / "saida" / f"Matriz_{defin['anexo'].replace(' ', '_')}_{slug(cfg['municipio']).replace('-', '_')}.xlsx"
    saida.parent.mkdir(exist_ok=True)
    wb.save(saida)
    print(f"Matriz: {saida}")
    print(f"  {tot['ATENDIDO']} atendidos, {tot['PARCIALMENTE ATENDIDO']} parciais, {tot['NÃO ATENDIDO']} não atendidos "
          f"(de {len(resultados)} itens)")
    for r in resultados:
        if r["situacao"] == "NÃO ATENDIDO":
            print(f"  NÃO ATENDIDO {r['item']}: {r['exigencia']} – {r['obs']}")
    return saida


def main():
    console_utf8()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("chave")
    a = ap.parse_args()
    gerar(projeto(a.chave))


if __name__ == "__main__":
    main()
